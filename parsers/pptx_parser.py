"""PowerPoint document parser using python-pptx."""

import io
import logging
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from dataclasses import dataclass

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from PIL import Image
except ImportError as e:
    raise ImportError(
        f"Required libraries not installed: {e}. "
        "Run: pip install python-pptx Pillow"
    )

logger = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    """Container for parsed document content."""
    text: str
    pages: List[Dict]  # Slides in PPTX context
    images: List[Dict]
    metadata: Dict
    word_count: int


class PPTXParser:
    """
    Extracts text and images from PowerPoint presentations.

    Uses python-pptx library to parse PPTX files.
    """

    def __init__(self, extract_images: bool = True):
        """
        Initialize PPTX parser.

        Args:
            extract_images: Extract images from slides for analysis
        """
        self.extract_images = extract_images

    def parse(self, file_path: Path) -> ParsedDocument:
        """
        Parse PPTX document and extract text and images.

        Args:
            file_path: Path to PPTX file

        Returns:
            ParsedDocument with extracted content

        Raises:
            FileNotFoundError: If file doesn't exist
            ValueError: If file is not a valid PPTX
        """
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        logger.info(f"Parsing PPTX: {file_path}")

        try:
            prs = Presentation(file_path)
        except Exception as e:
            raise ValueError(f"Failed to open PPTX file: {e}")

        slides_data = []
        full_text = []
        all_images = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_text = []
            slide_images = []

            # Extract text from all shapes
            for shape in slide.shapes:
                # Text frames
                if hasattr(shape, "text") and shape.text:
                    slide_text.append(shape.text)

                # Tables
                if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                    table_text = self._extract_table_text(shape)
                    if table_text:
                        slide_text.append(table_text)

                # Images
                if self.extract_images and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image_info = self._extract_image(shape, slide_num, len(slide_images))
                        if image_info:
                            slide_images.append(image_info)
                            all_images.append(image_info)
                    except Exception as e:
                        logger.debug(f"Failed to extract image from slide {slide_num}: {e}")

            # Notes (speaker notes)
            if slide.has_notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text
                if notes_text.strip():
                    slide_text.append(f"[Notes: {notes_text}]")

            page_text = "\n".join(slide_text)
            full_text.append(page_text)

            slides_data.append({
                "page_number": slide_num,
                "text": page_text,
                "word_count": len(page_text.split()),
                "image_count": len(slide_images)
            })

        combined_text = "\n\n".join(full_text)
        word_count = len(combined_text.split())

        # Get metadata
        metadata = self._extract_metadata(prs, file_path)

        logger.info(
            f"Extracted {len(slides_data)} slides, {word_count} words, "
            f"{len(all_images)} images from {file_path.name}"
        )

        return ParsedDocument(
            text=combined_text,
            pages=slides_data,
            images=all_images,
            metadata=metadata,
            word_count=word_count
        )

    def _extract_table_text(self, table_shape) -> str:
        """Extract text from a table shape."""
        rows = []
        table = table_shape.table

        for row in table.rows:
            cells = []
            for cell in row.cells:
                cells.append(cell.text.strip())
            rows.append(" | ".join(cells))

        return "\n".join(rows)

    def _extract_image(self, shape, slide_num: int, img_index: int) -> Optional[Dict]:
        """Extract image from a picture shape."""
        try:
            image = shape.image
            image_bytes = image.blob

            # Convert to PIL Image for size info
            pil_image = Image.open(io.BytesIO(image_bytes))

            return {
                "page": slide_num,
                "index": img_index,
                "format": image.ext,
                "size": pil_image.size,
                "bytes": image_bytes,
                "description": None  # To be filled by vision model
            }
        except Exception as e:
            logger.debug(f"Failed to extract image: {e}")
            return None

    def _extract_metadata(self, prs: Presentation, file_path: Path) -> Dict:
        """Extract PPTX metadata."""
        metadata = {
            "filename": file_path.name,
            "file_size_bytes": file_path.stat().st_size,
            "slide_count": len(prs.slides)
        }

        try:
            core_props = prs.core_properties

            metadata.update({
                "title": core_props.title or "",
                "author": core_props.author or "",
                "subject": core_props.subject or "",
                "created": core_props.created.isoformat() if core_props.created else None,
                "modified": core_props.modified.isoformat() if core_props.modified else None
            })
        except Exception as e:
            logger.debug(f"Failed to extract metadata: {e}")

        return metadata

    def validate_document(self, parsed_doc: ParsedDocument, min_words: int = 100) -> Tuple[bool, Optional[str]]:
        """
        Validate that document has sufficient content.

        Args:
            parsed_doc: Parsed document to validate
            min_words: Minimum word count required

        Returns:
            Tuple of (is_valid, error_message)
        """
        if parsed_doc.word_count < min_words:
            return False, (
                f"Insufficient content: only {parsed_doc.word_count} words extracted "
                f"(minimum: {min_words}). The presentation may be mostly images."
            )

        if not parsed_doc.text.strip():
            return False, "No text content extracted from presentation."

        return True, None
